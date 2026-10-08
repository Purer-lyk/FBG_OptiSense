"""Build Chinese report from completed, frozen-holdout dBm training artifacts."""
import json
from pathlib import Path
import zipfile
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm,Pt,Inches
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
OUT=ROOT/'outputs/model_comparisons/sensor_token26_dbm_three_sessions_random_v1'


def main():
    result=json.loads((OUT/'comparison.json').read_text(encoding='utf-8'))
    audit=json.loads((OUT/'data_audit.json').read_text(encoding='utf-8'))
    split=json.loads((OUT/'split.json').read_text(encoding='utf-8'))
    report=result['selected_report']; config=report['configuration']; test=report['metrics']['test']
    figures=OUT/'figures'; figures.mkdir(exist_ok=True)
    plt.rcParams.update({'font.sans-serif':['Microsoft YaHei','SimHei','DejaVu Sans'],'axes.unicode_minus':False})
    fig,ax=plt.subplots(figsize=(10,4))
    for trial in result['trials']:
        history=trial['history']; ax.plot([h['epoch'] for h in history],[h['validation_accuracy']*100 for h in history],label=trial['configuration']['id'])
    ax.axhline(80,color='gray',linestyle='--',label='80%目标'); ax.set(xlabel='训练轮次',ylabel='验证集位置准确率 / %',ylim=(0,102)); ax.legend(fontsize=8); ax.grid(alpha=.2)
    fig.tight_layout(); fig.savefig(figures/'validation_search.png',dpi=160); plt.close(fig)
    cm=np.asarray(test['confusion_matrix']); fig,ax=plt.subplots(figsize=(9,8)); im=ax.imshow(cm,cmap='Blues'); fig.colorbar(im,ax=ax)
    ax.set(xticks=range(26),yticks=range(26),xlabel='预测位置',ylabel='真实位置',title='选中模型随机帧测试集混淆矩阵')
    for a in range(26):
        for b in range(26):
            if cm[a,b]: ax.text(b,a,str(cm[a,b]),ha='center',va='center',fontsize=6,color='white' if cm[a,b]>cm.max()/2 else 'black')
    fig.tight_layout(); fig.savefig(figures/'test_confusion.png',dpi=160); plt.close(fig)
    pred=np.load(OUT/'test_predictions.npz'); fig,axes=plt.subplots(1,3,figsize=(12,4))
    for j,axis in enumerate(['Fx','Fy','Fz']):
        actual=pred['forces'][:,j]; predicted=pred['predicted_forces'][:,j]; valid=pred['force_valid']; low=min(actual[valid].min(),predicted[valid].min()); high=max(actual[valid].max(),predicted[valid].max())
        for sid in range(3):
            mask=valid&(pred['session']==sid); axes[j].scatter(actual[mask],predicted[mask],s=9,alpha=.6,label=f'会话{sid+1}')
        axes[j].plot([low,high],[low,high],color='gray',linestyle='--'); axes[j].set(title=f'{axis} MAE={test["force_mae_n"][j]:.3f} N',xlabel='真实力 / N',ylabel='预测力 / N'); axes[j].grid(alpha=.2)
    axes[0].legend(fontsize=8); fig.tight_layout(); fig.savefig(figures/'force_scatter.png',dpi=160); plt.close(fig)
    doc=Document(); section=doc.sections[0]; section.top_margin=Cm(1.8); section.bottom_margin=Cm(1.8); section.left_margin=Cm(1.8); section.right_margin=Cm(1.8)
    normal=doc.styles['Normal']; normal.font.name='Microsoft YaHei'; normal.font.size=Pt(10); normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    section.header.paragraphs[0].text='三号机 FBG | 三会话dBm Sensor Token训练'
    section.footer.paragraphs[0].text='随机帧开发评估，不代表跨会话或独立按压泛化；deployment_ready=false'
    def p(text): doc.add_paragraph(str(text))
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
        for c,text in zip(t.rows[0].cells,headers): c.text=str(text)
        for row in rows:
            for c,text in zip(t.add_row().cells,row): c.text=str(text)
        return t
    doc.add_heading('三组数据合并随机划分\ndBm Sensor Token位置与三维力训练报告',0)
    p(f'最终测试位置准确率：{test["position_accuracy"]*100:.2f}%；80%目标：'+('达到。' if report['target_met'] else '未达到。'))
    p('本次保留九个五维真实光栅Token、共享注意力backbone、26类互斥位置分类与Fx/Fy/Fz回归双头。报告数值均来自真实训练产物，没有改写原始数据或删除备份。')
    doc.add_heading('1. 数据与预处理',1)
    table(['会话','帧数','有效dBm','类别'],[[Path(s['path']).parent.name,s['lines'],s['accepted'],','.join(map(str,sorted(map(int,s['label_counts']))))] for s in audit['sources']])
    p(f'合计{audit["total"]}帧，有效真实力{audit["valid_force"]}帧。输入读取quality.stable_dbm，缺失时读取training_views.stable_optical_power_dbm；不使用ADC作为网络特征，不使用电压、Makima曲线或拟合峰中心。')
    p('没有额外时间滤波或拟合。45个保存的ADC换算dBm按9×5排列，逐测量维度（每个光栅每个点分别统计）用训练集均值/标准差标准化，共45组参数；验证、测试和实时推理复用保存参数。真实三轴力同样只用训练集统计标准化，缺失力用mask屏蔽，不补零。')
    p(f'dBm为工程估算值，不是功率计溯源标定。三组跨阻档位一致：selector={audit["feedback_selector"]}。逐帧核验ADC转换与保存dBm的最大差值={audit["maximum_adc_to_saved_dbm_error"]:.3g} dB。保存模型的实时输入仍接收ADC，先转换成同口径dBm后才送入网络。')
    p('补充审计：3200个45维dBm向量无完全重复，会话内和跨会话均无完全重复；这不代表邻帧彼此独立。旧会话有7帧各含1个ADC零值，相应保存dBm按半LSB检测下限处理（约−41.4384 dBm），并非该点真实功率读数。本次保持这7帧及原保存值不变，没有因审计补充重新训练或修改测试结果。')
    p('所有帧稳定性及光学精度均未通过认证，数值完整不等于质量通过；空载力在后两个会话之间也有零偏差异。完整只读审计见结果包readonly_audit.json。')
    doc.add_heading('2. 合并打乱划分与测试规则',1)
    table(['集合','帧数','用途'],[['训练',len(split['train']),'拟合模型、输入/力标准化'],['验证',len(split['validation']),'选择网络配置、checkpoint与最佳轮次'],['测试',len(split['test']),'选择完成后评价选中模型']])
    p('按用户要求：三会话先合并，再分层随机打乱，70%训练/15%验证/15%测试。固定划分种子20261008，完整索引保存split.json；所有候选配置使用相同划分。搜索计划在训练前写入search_plan.json，按验证准确率选中，同分时按验证交叉熵选择；没有按测试成绩挑模型、改划分或追分重训。')
    p('重要限制：同一次按压的相邻帧可同时进入不同集合，因此分数可能偏乐观，不是跨会话或独立重复按压性能。0号现在来自后两个会话，第三会话覆盖0～25，但随机帧划分仍没有独立会话测试。功能冒烟运行仅核验读取/保存/预测接口，不用于配置选择。')
    table(['会话','训练','验证','测试'],[[Path(s['path']).parent.name]+[int(np.sum(np.load(OUT/'dataset.npz')['session'][split[k]]==i)) for k in ['train','validation','test']] for i,s in enumerate(audit['sources'])])
    doc.add_heading('3. 网络与搜索结果',1)
    p('五维Token线性嵌入+可学习光栅身份编码→多头自注意力→按九个真实光栅顺序展平→共享MLP→两个预测头。相比前次均值池化，本次有序展平保留各光栅的输出特征，未增加不存在的光栅。')
    table(['候选','维度/层数','初始LR','Dropout/权重衰减','力损失权重','最佳轮次','验证准确率'],[[r['configuration']['id'],f"{r['configuration']['width']}/{r['configuration']['layers']}",r['configuration']['lr'],f"{r['configuration']['dropout']}/{r['configuration']['weight_decay']}",r['configuration']['force_weight'],r['best_epoch'],f"{r['best_validation_accuracy']*100:.2f}%"] for r in result['trials']])
    table(['选中配置','取值'],[[k,config[k]] for k in ['id','seed','width','heads','layers','ffn','shared','dropout','lr','weight_decay','force_weight','batch_size','max_epochs','patience']]+[['实际训练轮数',report['epochs_completed']],['交付checkpoint轮次',report['best_epoch']],['优化器','AdamW'],['学习率调度','余弦退火到1e-5'],['梯度裁剪','L2范数1.0'],['损失','训练类别权重交叉熵 + 配置中的force_weight × 标准化力SmoothL1']])
    doc.add_picture(str(figures/'validation_search.png'),width=Inches(6.7))
    doc.add_heading('4. 选中模型指标',1)
    table(['集合','准确率','宏F1','向量MAE/N'],[[name,f"{m['position_accuracy']*100:.2f}%",f"{m['macro_f1_present_classes']:.4f}",f"{m['force_vector_mae_n']:.4f}"] for name,m in report['metrics'].items()])
    table(['测试指标','结果'],[['仅1～25号准确率',f"{test['pressed_only_accuracy']*100:.2f}%"],['0号精确率',f"{test['no_press_precision']*100:.2f}%"],['0号召回率',f"{test['no_press_recall']*100:.2f}%"],['三维向量MAE',f"{test['force_vector_mae_n']:.4f} N"]])
    table(['轴','MAE/N','RMSE/N','R²'],[[axis]+[f"{test[key][i]:.4f}" for key in ['force_mae_n','force_rmse_n','force_r2']] for i,axis in enumerate(['Fx','Fy','Fz'])])
    table(['会话测试子集','帧数','位置准确率'],[[Path(s['path']).parent.name,int(np.sum(pred['session']==i)),f"{np.mean(pred['labels'][pred['session']==i]==pred['predicted_labels'][pred['session']==i])*100:.2f}%"] for i,s in enumerate(audit['sources'])])
    doc.add_picture(str(figures/'test_confusion.png'),width=Inches(6.5)); doc.add_picture(str(figures/'force_scatter.png'),width=Inches(6.7))
    doc.add_heading('5. 分数解释、部署和复现',1)
    p('不能将本次分数相对前次61.36%的变化全部归因于dBm：同时改变了数据量、随机划分、标准化、池化和优化参数。要判断dBm自身收益，需在这同一划分和同一网络下训练ADC对照；本次没有做该对照，也不声称跨会话泛化已改善到同等水平。')
    p('原始三份frames.jsonl训练前后SHA256一致。checkpoint重载、保存dBm与实时ADC转换预测一致性已核验。GUI仍在独立Torch子进程推理，不在Qt父进程导入Torch。')
    p('模型路径：'+report['model_path']); p('复现：在项目Python目录用虚拟环境运行run_dbm_three_session_training.py --output <新的输出目录>，禁止覆盖现有结果。')
    path=OUT/'三会话dBm_SensorToken随机划分训练报告.docx'; doc.save(path)
    text=f"# 三会话dBm训练\n\n测试准确率：{test['position_accuracy']*100:.2f}%\n验证准确率：{report['metrics']['validation']['position_accuracy']*100:.2f}%\n模型：{report['model_path']}\n\n随机帧划分，非跨会话性能。原始数据未修改。具体指标及参数见Word和JSON。\n"
    (OUT/'README.md').write_text(text,encoding='utf-8')
    archive=OUT/'dbm_three_sessions_training_artifacts.zip'
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for f in OUT.rglob('*'):
            if f.is_file() and f!=archive: z.write(f,f.relative_to(OUT))
        for name in ['touch_dbm_training.py','run_dbm_three_session_training.py','build_dbm_training_report.py','dbm_training_subprocess_tests.py','temporary_optical_units.py']: z.write(ROOT/name,'code/'+name)
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
    check=Document(path); assert len(check.inline_shapes)==3
    assert f"{test['position_accuracy']*100:.2f}%" in '\n'.join(p.text for p in check.paragraphs)
    print(json.dumps({'docx':str(path.resolve()),'zip':str(archive.resolve()),'tables':len(check.tables),'images':len(check.inline_shapes),'test_accuracy':test['position_accuracy'],'config':config,'test_metrics':test},ensure_ascii=False,indent=2))

if __name__=='__main__': main()
